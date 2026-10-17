"""
Step 6: Reproduce Figure 1 of Card and Krueger (1994)

"Distribution of Starting Wage Rates": the percentage of stores in New Jersey
and Pennsylvania at each starting wage, in February 1992 (wave 1) and November
1992 (wave 2).

Reads input/public.csv directly and writes:
  1. output/figures/figure1_replication.png       -- two-panel bar chart
  2. output/tables/figure1_wage_distribution.csv  -- verifiable bin counts
  3. output/tables/figure1_wage_distribution.xlsx -- wide view + Notes sheet

Binning: the x-axis uses the paper's wage points $4.25, $4.35, $4.45, ... $5.55
(0.10 wide bins centred on those points). A wage is assigned to the nearest
point; exact halfway values ($4.30, $4.40, ..., $5.00) are rounded up, and any
wage at or above $5.60 is top-coded into the $5.55 bin. Percentages use all
stores in each state as the denominator (NJ = 331, PA = 79); stores with a
missing starting wage are not assigned to a wage bin but stay in the denominator.
"""
import numpy as np
import pandas as pd
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parent.parent
INPUT_FILE = ROOT / "input" / "public.csv"
OUTPUT_PNG = ROOT / "output" / "figures" / "figure1_replication.png"
OUTPUT_CSV = ROOT / "output" / "tables" / "figure1_wage_distribution.csv"
OUTPUT_XLSX = ROOT / "output" / "tables" / "figure1_wage_distribution.xlsx"

raw = pd.read_csv(INPUT_FILE)

# ---- Wage bins: integer cents to avoid floating-point edge cases ----
EDGES_CENTS = np.arange(420, 570, 10)   # 420, 430, ..., 560 (bin lower edges)
LABELS_CENTS = np.arange(425, 565, 10)  # 425, 435, ..., 555 (bin centres)
WAGE_LABELS = [f"${c / 100:.2f}" for c in LABELS_CENTS]


def assign_bin(cents):
    """Assign each wage (in cents) to a bin centre, round-half-up, top-coded."""
    idx = np.searchsorted(EDGES_CENTS, cents, side="right") - 1
    return LABELS_CENTS[np.clip(idx, 0, len(LABELS_CENTS) - 1)]


def distribution(wage_col, wave_label):
    """Return a long DataFrame of bin counts and percentages for one wave."""
    rows = []
    for state_val, state_name in [(1, "NJ"), (0, "PA")]:
        # Denominator: ALL stores in the state. Stores with a missing wage are
        # not assigned to a wage bin, but they stay in the denominator.
        denom = int((raw["STATE"] == state_val).sum())
        mask = (raw["STATE"] == state_val) & raw[wage_col].notna()
        wages = raw.loc[mask, wage_col]
        cents = np.round(wages * 100).astype(int)
        bins = assign_bin(cents)
        counts = pd.Series(bins).value_counts()
        for label in LABELS_CENTS:
            c = int(counts.get(label, 0))
            rows.append({
                "panel": wave_label,
                "state": state_name,
                "wage_point": f"${label / 100:.2f}",
                "count": c,
                "percent": round(100.0 * c / denom, 3) if denom else 0.0,
                "total": denom,
            })
    return pd.DataFrame(rows)


wave1 = distribution("WAGE_ST", "February 1992")
wave2 = distribution("WAGE_ST2", "November 1992")
dist = pd.concat([wave1, wave2], ignore_index=True)

# ---- Save the verifiable CSV ----
OUTPUT_CSV.parent.mkdir(parents=True, exist_ok=True)
dist.to_csv(OUTPUT_CSV, index=False, encoding="utf-8-sig")

# ---- Wide Excel (one sheet per wave) + Notes ----
wb = Workbook()
for i, (wave_label, df_wave) in enumerate([("February 1992", wave1),
                                           ("November 1992", wave2)]):
    ws = wb.active if i == 0 else wb.create_sheet()
    ws.title = wave_label
    ws.cell(row=1, column=1, value=wave_label).font = Font(bold=True, size=12)
    headers = ["Wage point", "NJ count", "NJ %", "PA count", "PA %"]
    for j, h in enumerate(headers, start=1):
        ws.cell(row=2, column=j, value=h).font = Font(bold=True)
    r = 3
    for _, grp in df_wave.groupby("wage_point", sort=False):
        nj = grp[grp["state"] == "NJ"].iloc[0]
        pa = grp[grp["state"] == "PA"].iloc[0]
        for j, v in enumerate([grp["wage_point"].iloc[0], nj["count"],
                               nj["percent"], pa["count"], pa["percent"]], start=1):
            ws.cell(row=r, column=j, value=v)
        r += 1
    r += 1
    nj_total = df_wave[df_wave["state"] == "NJ"]["total"].iloc[0]
    pa_total = df_wave[df_wave["state"] == "PA"]["total"].iloc[0]
    ws.cell(row=r, column=1,
            value=f"Denominator (all stores in state): NJ = {nj_total}, PA = {pa_total}")

nws = wb.create_sheet("Notes")
notes = [
    "Figure 1 replication -- notes",
    "",
    "1. Wage points on the x-axis: $4.25, $4.35, $4.45, ..., $5.55 (0.10-wide bins centred on these points).",
    "2. Binning rule: each starting wage is assigned to the nearest wage point. Exact halfway values",
    "   ($4.30, $4.40, ..., $5.00, $5.10) are rounded up to the higher point; any wage at or above $5.60",
    "   is top-coded into the $5.55 bin. Nothing is silently dropped.",
    "3. 'Other' wage values found in the data and how they are binned (round to nearest point):",
    "   wave 1 -- 4.32->4.35, 4.37->4.35, 4.39->4.35, 4.40->4.45, 4.62->4.65, 4.67->4.65, 4.87->4.85,",
    "   5.06->5.05, 5.12->5.15, 5.37->5.35, 5.42->5.45, 5.56->5.55, 5.62->5.55 (top), 5.75->5.55 (top);",
    "   wave 2 -- 4.30->4.35, 4.38->4.35, 4.40->4.45, 4.90->4.95, 4.91->4.95, 5.14->5.15, 5.20->5.25,",
    "   5.28->5.25, 5.40->5.45, 5.67->5.55 (top), 5.75->5.55 (top), 6.25->5.55 (top).",
    "4. Denominator: Percentages use all stores in each state as the denominator (NJ = 331, PA = 79),",
    "   matching the paper. Stores with a missing starting wage are not assigned to a wage bin but",
    "   remain in the denominator, so the percentages do not sum to 100% within a panel.",
    "5. The public data has one more NJ wave-2 store at exactly $5.05 than the paper's data",
    "   (283 vs 282), so the NJ $5.05 share is slightly above the paper's 85.2%.",
]
for i, line in enumerate(notes, start=1):
    nws.cell(row=i, column=1, value=line)

wb.save(OUTPUT_XLSX)

# ---- Plot: two stacked panels ----
fig, axes = plt.subplots(2, 1, figsize=(9.5, 10), sharex=True)
x = np.arange(len(WAGE_LABELS))
width = 0.38

panels = [
    (axes[0], "February 1992", wave1, (0, 35)),
    (axes[1], "November 1992", wave2, (0, 90)),
]

for ax, wave_label, df_wave, ylim in panels:
    nj = df_wave[df_wave["state"] == "NJ"]["percent"].to_numpy()
    pa = df_wave[df_wave["state"] == "PA"]["percent"].to_numpy()
    ax.bar(x - width / 2, nj, width, label="New Jersey",
           color="#1f1f1f", edgecolor="#1f1f1f", zorder=3)
    ax.bar(x + width / 2, pa, width, label="Pennsylvania",
           color="white", edgecolor="#1f1f1f", hatch="////", linewidth=1.0, zorder=3)

    ax.set_title(wave_label, fontsize=14, fontweight="bold")
    ax.set_ylabel("Percent of stores", fontsize=12)
    ax.set_ylim(*ylim)
    ax.yaxis.set_major_locator(plt.MultipleLocator(5 if ylim[1] <= 40 else 15))
    ax.tick_params(axis="y", labelsize=10)
    ax.grid(axis="y", color="#d9d9d9", linewidth=0.8, zorder=0)
    ax.set_axisbelow(True)
    for spine in ["top", "right"]:
        ax.spines[spine].set_visible(False)

axes[1].set_xlabel("Starting wage ($/hour)", fontsize=12)
axes[1].set_xticks(x)
axes[1].set_xticklabels(WAGE_LABELS, rotation=90, fontsize=10)

handles, labels = axes[0].get_legend_handles_labels()
fig.legend(handles, labels, loc="upper center", ncol=2, frameon=False, fontsize=12,
           bbox_to_anchor=(0.5, 1.005))
fig.suptitle("Figure 1 -- Distribution of Starting Wage Rates "
             "(Card & Krueger 1994, replication)", fontsize=15, fontweight="bold", y=1.03)

fig.tight_layout(rect=[0, 0, 1, 0.97])
OUTPUT_PNG.parent.mkdir(parents=True, exist_ok=True)
fig.savefig(OUTPUT_PNG, dpi=200, bbox_inches="tight")
plt.close(fig)

# ---- Terminal summary: key percentages vs the paper ----
def key_pct(df_wave, state, point_cents):
    row = df_wave[(df_wave["state"] == state)
                  & (df_wave["wage_point"] == f"${point_cents / 100:.2f}")]
    return row["percent"].iloc[0], int(row["count"].iloc[0]), int(row["total"].iloc[0])


print("=== Step 6: reproduce Figure 1 ===")
print(f"Read {len(raw)} rows from input/public.csv.")
for wave_label, df_wave in [("February 1992", wave1), ("November 1992", wave2)]:
    nj_n = df_wave[df_wave["state"] == "NJ"]["total"].iloc[0]
    pa_n = df_wave[df_wave["state"] == "PA"]["total"].iloc[0]
    print(f"{wave_label}: denominator (all stores in state) -- NJ = {nj_n}, PA = {pa_n}")

print()
print("Key percentages (all-stores denominator):")
for wave_label, df_wave, point, state, paper_val in [
    ("February 1992", wave1, 425, "NJ", 30.5),
    ("November 1992", wave2, 425, "PA", 25.3),
    ("November 1992", wave2, 505, "NJ", 85.2),
]:
    pct, cnt, tot = key_pct(df_wave, state, point)
    print(f"  {wave_label:14s} {state} ${point / 100:.2f}: {pct:.1f}% "
          f"({cnt}/{tot})  vs paper {paper_val}%")

print()
print(f"Saved {OUTPUT_PNG}")
print(f"Saved {OUTPUT_CSV}")
print(f"Saved {OUTPUT_XLSX}")
