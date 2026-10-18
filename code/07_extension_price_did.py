"""
Step 7 (extension): Did New Jersey restaurants raise full-meal prices relative
to Pennsylvania after the minimum-wage increase?

This is a project extension, not a replication of the paper's Table 2-4 or
Figure 1. Full-meal price = medium soda + small fries + entree (each including
tax, in USD). The sample is the stores with non-missing full-meal prices in both
waves; closed stores and stores with any missing price are excluded (no price is
imputed to 0).

Reads data/raw/public.csv and writes:
  1. outputs/tables/extension_price_did.csv  -- means, DiD, and regressions
  2. outputs/tables/extension_price_did.xlsx -- results + Notes sheet
  3. outputs/figures/extension_price_did.png -- before/after mean-price lines
"""
import numpy as np
import pandas as pd
from pathlib import Path
import statsmodels.api as sm
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

ROOT = Path(__file__).resolve().parent.parent
INPUT_FILE = ROOT / "data" / "raw" / "public.csv"
OUTPUT_CSV = ROOT / "outputs" / "tables" / "extension_price_did.csv"
OUTPUT_XLSX = ROOT / "outputs" / "tables" / "extension_price_did.xlsx"
OUTPUT_PNG = ROOT / "outputs" / "figures" / "extension_price_did.png"

raw = pd.read_csv(INPUT_FILE)

# ---- Full-meal price (medium soda + small fries + entree, incl. tax, USD) ----
raw["meal_before"] = raw["PSODA"] + raw["PFRY"] + raw["PENTREE"]
raw["meal_after"] = raw["PSODA2"] + raw["PFRY2"] + raw["PENTREE2"]

# ---- Sample: both waves non-missing (no imputation for closed/missing) ----
sample_mask = raw["meal_before"].notna() & raw["meal_after"].notna()
d = raw.loc[sample_mask].copy()
d["nj"] = (d["STATE"] == 1).astype(int)
d["price_change"] = d["meal_after"] - d["meal_before"]

N = len(d)
nj_mask = d["nj"] == 1
pa_mask = d["nj"] == 0


def mean_se(series):
    x = series.dropna()
    n = len(x)
    return float(x.mean()), float(x.std(ddof=1) / np.sqrt(n)) if n > 1 else float("nan"), n


nj_before, nj_before_se, nj_n = mean_se(d.loc[nj_mask, "meal_before"])
nj_after, nj_after_se, _ = mean_se(d.loc[nj_mask, "meal_after"])
pa_before, pa_before_se, pa_n = mean_se(d.loc[pa_mask, "meal_before"])
pa_after, pa_after_se, _ = mean_se(d.loc[pa_mask, "meal_after"])

nj_change = nj_after - nj_before
pa_change = pa_after - pa_before
did = nj_change - pa_change

# ---- Simple OLS: price_change = constant + NJ dummy ----
X1 = sm.add_constant(d[["nj"]])
res1 = sm.OLS(d["price_change"], X1).fit()

# ---- Adjusted OLS: + chain dummies + company-owned dummy ----
chain = pd.DataFrame({f"chain_{c}": (d["CHAIN"] == c).astype(int) for c in [2, 3, 4]})
co_owned = pd.DataFrame({"co_owned": (d["CO_OWNED"] == 1).astype(int)})
X2 = sm.add_constant(pd.concat([d[["nj"]], chain, co_owned], axis=1))
res2 = sm.OLS(d["price_change"], X2).fit()


def ols_row(res, model):
    return {
        "section": "regression",
        "item": f"NJ dummy ({model})",
        "state": "", "period": "",
        "value": float(res.params["nj"]),
        "se": float(res.bse["nj"]),
        "t": float(res.tvalues["nj"]),
        "p_value": float(res.pvalues["nj"]),
        "n": N,
    }


records = [
    {"section": "means", "item": "mean full-meal price", "state": "NJ", "period": "before",
     "value": nj_before, "se": nj_before_se, "t": None, "p_value": None, "n": nj_n},
    {"section": "means", "item": "mean full-meal price", "state": "NJ", "period": "after",
     "value": nj_after, "se": nj_after_se, "t": None, "p_value": None, "n": nj_n},
    {"section": "means", "item": "mean full-meal price", "state": "PA", "period": "before",
     "value": pa_before, "se": pa_before_se, "t": None, "p_value": None, "n": pa_n},
    {"section": "means", "item": "mean full-meal price", "state": "PA", "period": "after",
     "value": pa_after, "se": pa_after_se, "t": None, "p_value": None, "n": pa_n},
    {"section": "change", "item": "price change", "state": "NJ", "period": "",
     "value": nj_change, "se": None, "t": None, "p_value": None, "n": nj_n},
    {"section": "change", "item": "price change", "state": "PA", "period": "",
     "value": pa_change, "se": None, "t": None, "p_value": None, "n": pa_n},
    {"section": "did", "item": "DiD (NJ - PA)", "state": "", "period": "",
     "value": did, "se": None, "t": None, "p_value": None, "n": N},
    ols_row(res1, "simple"),
    ols_row(res2, "adjusted"),
]

table = pd.DataFrame(records)
for col in ["value", "se", "t", "p_value"]:
    table[col] = table[col].round(4)

OUTPUT_CSV.parent.mkdir(parents=True, exist_ok=True)
table.to_csv(OUTPUT_CSV, index=False, encoding="utf-8-sig")

# ---- Excel ----
wb = Workbook()
ws = wb.active
ws.title = "Results"
ws.cell(row=1, column=1,
        value="Extension -- Full-Meal Price DiD: New Jersey vs Pennsylvania "
              "(minimum-wage increase)").font = Font(bold=True, size=12)

headers = ["Section", "Item", "State", "Period", "Value", "SE", "t", "p-value", "n"]
for j, h in enumerate(headers, start=1):
    ws.cell(row=2, column=j, value=h).font = Font(bold=True)

r = 3
for rec in records:
    ws.cell(row=r, column=1, value=rec["section"])
    ws.cell(row=r, column=2, value=rec["item"])
    ws.cell(row=r, column=3, value=rec["state"])
    ws.cell(row=r, column=4, value=rec["period"])
    ws.cell(row=r, column=5, value=round(rec["value"], 4))
    ws.cell(row=r, column=6, value=round(rec["se"], 4) if rec["se"] is not None else "")
    ws.cell(row=r, column=7, value=round(rec["t"], 4) if rec["t"] is not None else "")
    ws.cell(row=r, column=8, value=round(rec["p_value"], 4) if rec["p_value"] is not None else "")
    ws.cell(row=r, column=9, value=rec["n"])
    r += 1

nws = wb.create_sheet("Notes")
notes = [
    "Extension analysis -- notes",
    "",
    "This is a PROJECT EXTENSION, not a replication of the paper's Table 2-4 or Figure 1.",
    "",
    "1. Research question: Did New Jersey restaurants raise full-meal prices relative to",
    "   Pennsylvania after the April 1992 minimum-wage increase ($4.25 -> $5.05 in NJ only)?",
    "2. Full-meal price = medium soda (PSODA/PSODA2) + small fries (PFRY/PFRY2) + entree",
    "   (PENTREE/PENTREE2), each including tax, in USD. Same definition as Table 2's 'price of full meal'.",
    "3. Sample: stores with non-missing full-meal price in BOTH waves (n = 356: NJ 288, PA 68).",
    "   Closed stores and stores with any missing price component are excluded; no price is imputed to 0.",
    "4. DiD = (NJ after - NJ before) - (PA after - PA before). The simple OLS of price_change on the",
    "   NJ dummy returns exactly this DiD as the NJ coefficient; the adjusted model adds three chain",
    "   dummies (Burger King omitted) and a company-owned dummy as a robustness check.",
    "5. Interpretation: a positive DiD means NJ prices rose relative to PA. This is consistent with",
    "   some pass-through of the higher NJ minimum wage into menu prices, but the estimate is small",
    "   and the extension is observational; see the main report for the significance and caveats.",
]
for i, line in enumerate(notes, start=1):
    nws.cell(row=i, column=1, value=line)

wb.save(OUTPUT_XLSX)

# ---- Figure: before/after mean full-meal price by state ----
fig, ax = plt.subplots(figsize=(8, 6))
x = [0, 1]
tick_labels = ["Before\n(Feb 1992)", "After\n(Nov 1992)"]

ax.plot(x, [nj_before, nj_after], marker="o", markersize=9, linewidth=2.2,
        color="#1f1f1f", label="New Jersey", zorder=3)
ax.plot(x, [pa_before, pa_after], marker="s", markersize=9, linewidth=2.2,
        color="#8a8a8a", linestyle="--", label="Pennsylvania", zorder=3)

for xpos, yval in [(0, nj_before), (1, nj_after), (0, pa_before), (1, pa_after)]:
    ax.annotate(f"{yval:.2f}", (xpos, yval), textcoords="offset points",
                xytext=(0, 10 if "NJ" else 10), ha="center", fontsize=10)

ax.set_xticks(x)
ax.set_xticklabels(tick_labels, fontsize=11)
ax.set_ylabel("Average full-meal price (USD)", fontsize=12)
ax.set_ylim(2.95, 3.55)
ax.set_xlim(-0.3, 1.3)
ax.grid(axis="y", color="#d9d9d9", linewidth=0.8, zorder=0)
ax.set_axisbelow(True)
for spine in ["top", "right"]:
    ax.spines[spine].set_visible(False)

p_simple = res1.pvalues["nj"]
ax.set_title("Full-Meal Price Before and After the Minimum-Wage Increase",
             fontsize=13, fontweight="bold")
ax.text(0.5, 3.02, f"DiD (NJ - PA) = {did:+.3f} USD   (OLS p = {p_simple:.3f})",
        ha="center", fontsize=11)
ax.legend(loc="lower right", frameon=False, fontsize=11)

fig.tight_layout()
OUTPUT_PNG.parent.mkdir(parents=True, exist_ok=True)
fig.savefig(OUTPUT_PNG, dpi=200, bbox_inches="tight")
plt.close(fig)

# ---- Terminal summary ----
print("=== Step 7 (extension): full-meal price DiD ===")
print(f"Sample (both waves non-missing full-meal price): n = {N} (NJ {nj_n}, PA {pa_n})")
print()
print(f"NJ  before = {nj_before:.3f}, after = {nj_after:.3f}, change = {nj_change:+.3f}")
print(f"PA  before = {pa_before:.3f}, after = {pa_after:.3f}, change = {pa_change:+.3f}")
print(f"DiD (NJ - PA) = {did:+.3f}")
print()
print("Simple OLS (price_change ~ NJ dummy):")
print(f"  coef = {res1.params['nj']:+.4f}, SE = {res1.bse['nj']:.4f}, "
      f"t = {res1.tvalues['nj']:.3f}, p = {res1.pvalues['nj']:.4f}, n = {N}")
print("Adjusted OLS (chain + company-owned controls):")
print(f"  coef = {res2.params['nj']:+.4f}, SE = {res2.bse['nj']:.4f}, "
      f"t = {res2.tvalues['nj']:.3f}, p = {res2.pvalues['nj']:.4f}, n = {N}")
print()
print(f"Saved {OUTPUT_CSV}")
print(f"Saved {OUTPUT_XLSX}")
print(f"Saved {OUTPUT_PNG}")
